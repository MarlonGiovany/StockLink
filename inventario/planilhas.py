"""Relatórios em Excel (.xlsx).

Cada função devolve a pasta de trabalho já montada; a view só entrega o
arquivo. Os valores vão como número de verdade (não texto), para quem abrir
no Excel poder somar e filtrar sem converter nada.
"""

import re
from io import BytesIO

from django.db.models import Count, Prefetch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Locacao

FORMATO_REAIS = 'R$ #,##0.00'
FORMATO_DATA = "DD/MM/YYYY"


def _planilha(titulo, cabecalho):
    livro = Workbook()
    folha = livro.active
    folha.title = titulo
    folha.append(cabecalho)
    for celula in folha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="162D63")
    # Cabeçalho fica parado no topo ao rolar
    folha.freeze_panes = "A2"
    return livro, folha


def _acerta_colunas(folha):
    for coluna in folha.columns:
        maior = max(len(str(c.value)) if c.value is not None else 0 for c in coluna)
        folha.column_dimensions[get_column_letter(coluna[0].column)].width = min(maior + 3, 50)
    folha.auto_filter.ref = folha.dimensions


def _formata(folha, coluna, formato):
    for (celula,) in folha.iter_rows(min_row=2, min_col=coluna, max_col=coluna):
        celula.number_format = formato


def em_bytes(livro):
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def maquinas_locadas(cliente=None, com_contrato=False):
    """Uma linha por máquina em locação ativa, agrupada por cliente.

    O número do contrato só entra para quem pode ver a aba Contratos.
    """
    cabecalho = ["Cliente", "CPF/CNPJ", "INF", "Produto", "Marca", "Modelo",
                 "Nº de série", "Local", "Início da locação", "Valor da locação"]
    if com_contrato:
        cabecalho.append("Contrato")
    livro, folha = _planilha("Máquinas locadas", cabecalho)

    locacoes = (
        Locacao.objects.filter(ativa=True)
        .select_related("cliente", "equipamento", "equipamento__produto", "contrato")
        .order_by("cliente__nome", "equipamento__numero_patrimonio")
    )
    if cliente is not None:
        locacoes = locacoes.filter(cliente=cliente)

    for loc in locacoes:
        m = loc.equipamento
        linha = [
            loc.cliente.nome, loc.cliente.documento, f"INF-{m.numero_patrimonio}",
            m.produto.nome if m.produto else "", m.marca, m.modelo,
            m.numero_serie, m.local, loc.data_inicio, loc.valor,
        ]
        if com_contrato:
            linha.append(loc.contrato.numero if loc.contrato else "")
        folha.append(linha)

    _formata(folha, 9, FORMATO_DATA)
    _formata(folha, 10, FORMATO_REAIS)
    _acerta_colunas(folha)
    return livro


def equipamentos(maquinas, titulo, com_locacao=False, com_contrato=False):
    """Os equipamentos do filtro da tela, um por linha.

    `com_locacao` acrescenta cliente, início e valor da locação ativa — só
    faz sentido para as locadas (nas outras situações essas colunas sairiam
    vazias). A ficha técnica (processador, memória, armazenamento) só entra
    se alguma máquina da lista tiver ela preenchida.
    """
    maquinas = list(
        maquinas.select_related("produto", "fornecedor")
        # distinct: o filtro por cliente junta as locações e multiplicaria a conta
        .annotate(qtd_manutencoes=Count("manutencoes", distinct=True))
        .prefetch_related(Prefetch(
            "locacoes",
            queryset=Locacao.objects.filter(ativa=True).select_related("cliente", "contrato"),
            to_attr="locacao_ativa_lista",
        ))
        .order_by("produto__nome", "numero_patrimonio")
    )
    com_ficha = any(m.processador or m.memoria_ram or m.armazenamento for m in maquinas)

    colunas = [
        ("INF", lambda m, l: f"INF-{m.numero_patrimonio}", None),
        ("Produto", lambda m, l: m.produto.nome if m.produto else "", None),
        ("Marca", lambda m, l: m.marca, None),
        ("Modelo", lambda m, l: m.modelo, None),
        ("Nº de série", lambda m, l: m.numero_serie, None),
        ("Status", lambda m, l: m.get_status_display(), None),
        ("Local", lambda m, l: m.local, None),
    ]
    if com_locacao:
        colunas += [
            ("Cliente", lambda m, l: l.cliente.nome if l else "", None),
            ("CPF/CNPJ", lambda m, l: l.cliente.documento if l else "", None),
            ("Início da locação", lambda m, l: l.data_inicio if l else None, FORMATO_DATA),
            ("Valor da locação", lambda m, l: l.valor if l else None, FORMATO_REAIS),
        ]
        if com_contrato:
            colunas.append(
                ("Contrato", lambda m, l: l.contrato.numero if l and l.contrato else "", None)
            )
    if com_ficha:
        colunas += [
            ("Processador", lambda m, l: m.processador, None),
            ("Memória RAM", lambda m, l: m.memoria_ram, None),
            ("Armazenamento", lambda m, l: m.armazenamento, None),
        ]
    colunas += [
        ("Fornecedor", lambda m, l: m.fornecedor.nome if m.fornecedor else "", None),
        ("Data de aquisição", lambda m, l: m.data_aquisicao, FORMATO_DATA),
        ("Valor pago", lambda m, l: m.valor_compra, FORMATO_REAIS),
        ("Manutenções", lambda m, l: m.qtd_manutencoes, None),
        ("Observações", lambda m, l: m.observacoes, None),
    ]

    # O Excel recusa \ / ? * [ ] : no nome da aba ("Baixado / Descartado")
    # e corta em 31 letras
    aba = re.sub(r"[\\/?*\[\]:]", "-", titulo)[:31]
    livro, folha = _planilha(aba, [nome for nome, _, _ in colunas])
    for m in maquinas:
        loc = m.locacao_ativa_lista[0] if m.locacao_ativa_lista else None
        folha.append([valor(m, loc) for _, valor, _ in colunas])
    for n, (_, _, formato) in enumerate(colunas, start=1):
        if formato:
            _formata(folha, n, formato)
    _acerta_colunas(folha)
    return livro
