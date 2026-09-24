"""Importa clientes de uma planilha Excel (.xlsx), cadastrando só os novos.

Uso:
    python manage.py importar_clientes "C:\\caminho\\clientes.xlsx" --simular
    python manage.py importar_clientes "C:\\caminho\\clientes.xlsx"

A primeira linha da planilha deve ser o cabeçalho. Colunas reconhecidas
(maiúsculas e acentos não importam): nome, documento (CPF/CNPJ), telefone,
e-mail, endereço e observações. Só o nome é obrigatório.

Um cliente é considerado já cadastrado quando:
  - a linha tem CPF/CNPJ e já existe cliente com os mesmos dígitos (mesmo que o
    nome esteja escrito de outro jeito); ou
  - a linha não tem CPF/CNPJ e já existe cliente com o mesmo nome.

Mesmo nome com CPF/CNPJ diferente (ou com CPF/CNPJ quando o cadastrado não tem)
NÃO é cadastrado: a linha vai para a lista "para revisar" do relatório, porque
pode ser filial ou erro de digitação no CNPJ.

Com --filiais, o CPF/CNPJ sozinho deixa de bastar: filiais que dividem o mesmo
CNPJ entram como clientes novos, e só é duplicata quando nome e CNPJ batem.
Use só quando a planilha escreve os nomes do mesmo jeito que o sistema, senão
o mesmo cliente com outro nome é cadastrado de novo.
"""
import re
import unicodedata

from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from inventario.models import Cliente

COLUNAS = {
    "nome": ("nome", "razao social", "nome razao social", "cliente"),
    "documento": ("documento", "cpf", "cnpj", "cpnj", "cpf cnpj", "cpf/cnpj"),
    "telefone": ("telefone", "fone", "celular", "contato"),
    "email": ("email", "e-mail"),
    "endereco": ("endereco", "endereço"),
    "observacoes": ("observacoes", "observacao", "obs"),
}


def _sem_acento(texto):
    base = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in base if not unicodedata.combining(c))


def _chave_nome(nome):
    return " ".join(_sem_acento(nome).casefold().split())


def _chave_documento(documento):
    return re.sub(r"\D", "", documento)


def _sem_palavras(nome, palavras):
    """Tira palavras soltas (ou entre parênteses) do nome, ex.: 'NOTEBOOK'."""
    for palavra in palavras:
        p = re.escape(palavra)
        nome = re.sub(rf"\(\s*{p}\s*\)", " ", nome, flags=re.IGNORECASE)
        nome = re.sub(rf"\b{p}\b", " ", nome, flags=re.IGNORECASE)
    return " ".join(nome.split())


def _texto(valor):
    """Célula do Excel -> texto. Números inteiros não viram '123.0'."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


class Command(BaseCommand):
    help = "Importa clientes de uma planilha Excel, ignorando os já cadastrados."

    def add_arguments(self, parser):
        parser.add_argument("arquivo", type=str, help="Caminho do arquivo .xlsx")
        parser.add_argument(
            "--simular", action="store_true",
            help="Mostra o que seria feito, sem gravar nada no banco.",
        )
        parser.add_argument(
            "--sem-parenteses", action="store_true",
            help='Tira do nome o trecho final entre parênteses, ex.: "JAV MATRIZ (COD. X)".',
        )
        parser.add_argument(
            "--filiais", action="store_true",
            help="Mesmo CPF/CNPJ com nome diferente entra como cliente novo (filial).",
        )
        parser.add_argument(
            "--remover-palavras", type=str, default="",
            help='Palavras a tirar do nome, separadas por vírgula, ex.: "notebook,impressora".',
        )

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise CommandError("Falta instalar o openpyxl: pip install openpyxl")

        try:
            planilha = load_workbook(options["arquivo"], read_only=True, data_only=True)
        except (OSError, ValueError) as erro:
            raise CommandError(f"Não consegui abrir o arquivo: {erro}")

        try:
            linhas = iter(list(planilha.active.iter_rows(values_only=True)))
        finally:
            planilha.close()
        cabecalho = next(linhas, None)
        if not cabecalho:
            raise CommandError("A planilha está vazia.")
        indices = self._mapear_colunas(cabecalho)

        palavras = [p.strip() for p in options["remover_palavras"].split(",") if p.strip()]

        # nome (sem acento/caixa) -> CPF/CNPJs (só dígitos) já cadastrados com ele
        docs_por_nome, docs_existentes = {}, set()
        for nome, documento in Cliente.objects.values_list("nome", "documento"):
            docs_por_nome.setdefault(_chave_nome(nome), set()).add(_chave_documento(documento))
            docs_existentes.add(_chave_documento(documento))
        docs_existentes.discard("")

        criados, ignorados, para_revisar, com_erro = [], [], [], []
        with transaction.atomic():
            for numero, linha in enumerate(linhas, start=2):
                dados = {
                    campo: _texto(linha[i]) if i < len(linha) else ""
                    for campo, i in indices.items()
                }
                if not any(dados.values()):
                    continue
                if palavras and dados.get("nome"):
                    dados["nome"] = _sem_palavras(dados["nome"], palavras)
                if options["sem_parenteses"] and dados.get("nome"):
                    dados["nome"] = re.sub(r"\s*\([^)]*\)\s*$", "", dados["nome"])
                if not dados.get("nome"):
                    com_erro.append((numero, "sem nome"))
                    continue

                doc = _chave_documento(dados.get("documento", ""))
                if not doc:
                    dados["documento"] = ""  # ex.: "SEM PREENCHIMENTO"
                if doc and doc in docs_existentes and not options["filiais"]:
                    ignorados.append((numero, dados["nome"]))
                    continue
                chave = _chave_nome(dados["nome"])
                docs_do_nome = docs_por_nome.get(chave)
                if docs_do_nome is not None:
                    if not doc or doc in docs_do_nome:
                        ignorados.append((numero, dados["nome"]))
                        continue
                    outros = ", ".join(sorted(d for d in docs_do_nome if d)) or "nenhum"
                    para_revisar.append((
                        numero,
                        f'"{dados["nome"]}" ({dados["documento"]}) já existe com '
                        f"outro CPF/CNPJ (cadastrado: {outros})",
                    ))
                    continue

                cliente = Cliente(**dados)
                try:
                    cliente.full_clean()
                except ValidationError as erro:
                    detalhe = "; ".join(
                        f"{campo}: {' '.join(msgs)}"
                        for campo, msgs in erro.message_dict.items()
                    )
                    com_erro.append((numero, detalhe))
                    continue

                if not options["simular"]:
                    cliente.save()
                docs_por_nome.setdefault(chave, set()).add(doc)
                if doc:
                    docs_existentes.add(doc)
                criados.append((numero, cliente.nome))

        for numero, motivo in com_erro:
            self.stdout.write(self.style.WARNING(f"Linha {numero} não importada: {motivo}"))
        for numero, motivo in para_revisar:
            self.stdout.write(self.style.WARNING(f"Linha {numero} para revisar: {motivo}"))
        prefixo = "SIMULAÇÃO (nada foi gravado) - " if options["simular"] else ""
        self.stdout.write(self.style.SUCCESS(
            f"{prefixo}{len(criados)} novos, {len(ignorados)} já cadastrados, "
            f"{len(para_revisar)} para revisar, {len(com_erro)} com problema."
        ))

    def _mapear_colunas(self, cabecalho):
        indices = {}
        for i, titulo in enumerate(cabecalho):
            titulo = _chave_nome(_texto(titulo))
            for campo, apelidos in COLUNAS.items():
                if campo not in indices and titulo in {_chave_nome(a) for a in apelidos}:
                    indices[campo] = i
        if "nome" not in indices:
            raise CommandError(
                "Não achei a coluna de nome na primeira linha da planilha "
                "(esperado: Nome, Razão social ou Cliente)."
            )
        return indices
